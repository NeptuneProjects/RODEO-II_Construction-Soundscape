"""Word documents built from report content, for pasting into a shared report.

Tables are native Word tables and captions carry SEQ fields, so figures and
tables renumber in whatever document they are pasted into (select all and
press F9 to update fields).

Text may mark subscripts and superscripts as `L_{p,pk}` and `uPa^{2}`.
"""

import io
import re
from pathlib import Path
from typing import BinaryIO

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image

_MARKUP = re.compile(r"([_^])\{([^}]*)\}")


def add_text(paragraph: Paragraph, text: str, bold: bool = False) -> None:
    """Append text to a paragraph, rendering `_{...}` and `^{...}` as
    subscript and superscript."""
    pos = 0
    for m in _MARKUP.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()]).bold = bold
        run = paragraph.add_run(m.group(2))
        run.bold = bold
        if m.group(1) == "_":
            run.font.subscript = True
        else:
            run.font.superscript = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:]).bold = bold


def plain(text: str) -> str:
    """Text with the sub- and superscript markup removed."""
    return _MARKUP.sub(r"\2", text)


def _seq_field(paragraph: Paragraph, name: str, number: int) -> None:
    """A SEQ field showing `number` until Word updates it."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f" SEQ {name} \\* ARABIC ")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = str(number)
    run.append(t)
    field.append(run)
    paragraph._p.append(field)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _keep_together(table) -> None:
    """Move the whole table to the next page rather than split it, when it
    fits on one page: no row breaks across pages, and every row but the last
    is kept with the next. A table longer than a page still breaks, with its
    header row repeated."""
    for i, row in enumerate(table.rows):
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cant_split)
        if i < len(table.rows) - 1:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True


# Parts of a .docx that carry its styles and page setup rather than its content.
_STYLE_PARTS = {
    "styles",
    "stylesWithEffects",
    "numbering",
    "settings",
    "theme",
    "fontTable",
    "webSettings",
    "header",
    "footer",
    "footnotes",
    "endnotes",
}


def styles_only(source: Path, destination: Path | BinaryIO) -> None:
    """Copy a .docx keeping only its styles and page setup.

    The body is emptied except for the final section properties (page size,
    margins, headers, and footers), and every part the body referred to
    (images, charts, comments, hyperlinks, custom XML) is dropped, so the copy
    is small enough to serve as a style template. `source` is not modified.
    """
    doc = docx.Document(str(source))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    rels = doc.part.rels
    for r_id in [k for k, r in rels.items() if r.reltype.rsplit("/", 1)[-1] not in _STYLE_PARTS]:
        rels.pop(r_id)
    doc.save(destination if not isinstance(destination, Path) else str(destination))


class Document:
    """A Word document of headings, text, figures, and tables.

    `template` is a .docx whose styles (and page setup) are used; its content
    is discarded (see `styles_only`). Download the shared report as .docx to
    match its styles.
    """

    def __init__(self, template: Path | None = None):
        if template is None:
            self.doc = docx.Document()
        else:
            buffer = io.BytesIO()
            styles_only(template, buffer)
            buffer.seek(0)
            self.doc = docx.Document(buffer)
        # Word defines the Caption style in a file only once it is used, so a
        # template may lack it.
        styles = self.doc.styles
        if "Caption" not in [s.name for s in styles]:
            caption = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
            caption.base_style = styles["Normal"]
            caption.font.size = Pt(9)
            caption.font.italic = True
            caption.quick_style = True
        # Black captions, as in the documents pandoc writes (Word's default
        # Caption style is blue).
        styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
        self.counts = {"Figure": 0, "Table": 0}

    @property
    def _section(self):
        return self.doc.sections[-1]

    @property
    def text_width(self) -> Emu:
        s = self._section
        return Emu(s.page_width - s.left_margin - s.right_margin)

    @property
    def text_height(self) -> Emu:
        s = self._section
        return Emu(s.page_height - s.top_margin - s.bottom_margin)

    def orientation(self, landscape: bool) -> None:
        """Start a new section in landscape or portrait, if not already."""
        s = self._section
        if (s.orientation == WD_ORIENT.LANDSCAPE) == landscape:
            return
        if len(self.doc.paragraphs) == 0 and len(self.doc.tables) == 0:
            new = s  # Nothing written yet: change the first section.
        else:
            new = self.doc.add_section(WD_SECTION.NEW_PAGE)
        new.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        w, h = sorted((new.page_width, new.page_height))
        new.page_width, new.page_height = (h, w) if landscape else (w, h)

    def heading(self, text: str, level: int = 1) -> None:
        add_text(self.doc.add_heading(level=level), text)

    def paragraph(self, text: str) -> None:
        add_text(self.doc.add_paragraph(), text)

    def page_break(self) -> None:
        self.doc.add_page_break()

    def caption(self, kind: str, text: str) -> None:
        self.counts[kind] += 1
        p = self.doc.add_paragraph(style="Caption")
        p.add_run(f"{kind} ")
        _seq_field(p, kind, self.counts[kind])
        add_text(p, f": {text}")

    def figure(self, path: Path, caption: str, max_height: float = 0.85) -> None:
        """An image at the text width, or shorter if it would be taller than
        `max_height` of the text height, with its caption below."""
        with Image.open(path) as im:
            w, h = im.size
        width = min(self.text_width, int(max_height * self.text_height * w / h))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Emu(width))
        self.caption("Figure", caption)

    def table(
        self,
        columns: list[str],
        rows: list[list[str]],
        caption: str,
        align: str | None = None,
        font_pt: float = 8.0,
    ) -> None:
        """A table with its caption above, kept on one page where it fits
        (see `_keep_together`). `align` has one letter per column (l, c, r).
        A row's trailing empty cells are merged into the last nonempty one."""
        self.caption("Table", caption)
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        align = align or "l" * len(columns)
        just = {
            "l": WD_ALIGN_PARAGRAPH.LEFT,
            "c": WD_ALIGN_PARAGRAPH.CENTER,
            "r": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        table = self.doc.add_table(rows=1 + len(rows), cols=len(columns))
        table.style = "Table Grid"
        _repeat_header(table.rows[0])
        for i, values in enumerate([columns, *rows]):
            cells = table.rows[i].cells
            for j, value in enumerate(values):
                p = cells[j].paragraphs[0]
                p.alignment = just[align[j]] if i else WD_ALIGN_PARAGRAPH.CENTER
                add_text(p, value, bold=i == 0)
                for run in p.runs:
                    run.font.size = Pt(font_pt)
            if i:
                last = max(j for j, v in enumerate(values) if v or j == 0)
                if last < len(columns) - 1:
                    cells[last].merge(cells[-1])
        _keep_together(table)
        # Space after the table, so consecutive tables or figures separate.
        self.doc.add_paragraph()

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
