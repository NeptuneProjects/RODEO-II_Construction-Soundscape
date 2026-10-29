"""Correlation between pulses on a synthetic recording."""

import docx
import numpy as np
import polars as pl
import pytest

from soundscape import correlation, detect, extract, workflow
from soundscape.config import CorrelationConfig
from soundscape.correlation import (
    binned,
    decorrelation_time,
    histogram,
    level_label,
    read_correlation,
    stacked,
)
from tests.conftest import SyntheticRecording


def brute_corr_matrix(x: np.ndarray, threads: int = 1):
    """The maximum normalized cross-correlation and its lag (samples,
    positive when row i is later than row j), pair by pair with numpy."""
    z = (x - x.mean(-1, keepdims=True)) / x.std(-1, keepdims=True)
    n, m = z.shape
    corr, lag = np.ones((n, n)), np.zeros((n, n), int)
    for i in range(n):
        for j in range(n):
            c = np.correlate(z[i], z[j], "full") / np.sqrt(
                np.sum(z[i] ** 2) * np.sum(z[j] ** 2)
            )
            corr[i, j], lag[i, j] = c.max(), np.argmax(c) - (m - 1)
    return corr, lag


def fft_corr_matrix(x: np.ndarray, threads: int = 1):
    """The same as `brute_corr_matrix`, one row at a time with FFTs, fast
    enough for the hundreds of pulses in the synthetic recording."""
    z = (x - x.mean(-1, keepdims=True)) / x.std(-1, keepdims=True)
    n, m = z.shape
    n_fft = 2 * m
    spec = np.fft.rfft(z, n_fft)
    norm = np.sqrt(np.sum(z**2, axis=-1))
    lags = np.r_[np.arange(m), np.arange(-m, 0)]
    corr, lag = np.empty((n, n)), np.empty((n, n), int)
    for i in range(n):
        # irfft(X_i conj(X_j))[k] = sum_t z_i[t + k] z_j[t]
        c = np.fft.irfft(spec[i] * np.conj(spec), n_fft)
        k = np.argmax(c, axis=-1)
        corr[i] = c[np.arange(n), k] / (norm[i] * norm)
        lag[i] = lags[k]
    return corr, lag


@pytest.fixture
def run(site, monkeypatch):
    rec = SyntheticRecording(duration_s=900.0, gap=(300.5, 330.0))
    for module in (detect, extract):
        monkeypatch.setattr(module, "read_segment", rec.read)
    monkeypatch.setattr(correlation, "corr_matrix", fft_corr_matrix)
    # The synthetic pulse is a 200 Hz tone sampled at 2 kHz.
    site.correlation = CorrelationConfig(band=(50.0, 400.0), max_dt_s=60.0)
    site.min_strikes = 10
    w = workflow.Workflow(site)
    w.detect()
    w.extract()
    w.metrics()
    w.correlation()
    return rec


def test_julia_matches_numpy():
    rng = np.random.default_rng(1)
    x = rng.standard_normal((6, 300))
    x[1] = np.roll(x[0], 5)
    corr, lag = correlation.corr_matrix(x, threads=1)
    ref_corr, ref_lag = brute_corr_matrix(x)
    np.testing.assert_allclose(corr, ref_corr, atol=1e-9)
    np.testing.assert_array_equal(lag, ref_lag)
    assert lag[1, 0] == 5
    fft_corr, fft_lag = fft_corr_matrix(x)
    np.testing.assert_allclose(fft_corr, ref_corr, atol=1e-9)
    np.testing.assert_array_equal(fft_lag, ref_lag)


def test_condition_downsamples_within_band():
    config = CorrelationConfig(band=(10.0, 200.0), target_fs=1000.0)
    fs = 19531.25
    t = np.arange(round(1.2 * fs)) / fs
    w = np.sin(2 * np.pi * 50 * t) + np.sin(2 * np.pi * 2000 * t)
    y, fs_out = correlation.condition(w[None], fs, config)
    assert fs_out == pytest.approx(fs / 19)
    assert y.shape[-1] == len(t[::19])
    # The 50 Hz tone passes (away from the tapered ends); 2 kHz does not.
    mid = y[0, len(y[0]) // 4 : -len(y[0]) // 4]
    assert np.sqrt(np.mean(mid**2)) == pytest.approx(np.sqrt(0.5), rel=0.02)


def test_band_must_lie_below_nyquist():
    with pytest.raises(ValueError):
        CorrelationConfig(band=(10.0, 600.0), target_fs=1000.0)


def test_full_matrices(run, site):
    data, attrs = read_correlation(site.correlation_path("E1", "VLA1"))
    n = len(data["strike"])
    assert data["corr"].shape == data["lag"].shape == (2, n, n)
    assert attrs["sampling_rate"] == pytest.approx(1000.0)
    for c in range(2):
        corr = data["corr"][c]
        ok = np.isfinite(np.diagonal(corr))
        assert ok.sum() > 100
        np.testing.assert_allclose(np.diagonal(corr)[ok], 1.0, atol=1e-6)
        np.testing.assert_allclose(corr, corr.T, equal_nan=True)
        np.testing.assert_allclose(data["lag"][c], -data["lag"][c].T, equal_nan=True)
        # Invalid pulses (the gap) have NaN rows and columns.
        assert np.all(np.isnan(corr[~ok]))
        # Identical pulses at full amplitude, aligned at their detection
        # times.
        full = ok & (np.arange(n) > 100)
        sub = corr[np.ix_(full, full)]
        assert np.nanmin(sub) > 0.95
        assert np.nanmax(np.abs(data["lag"][c][np.ix_(full, full)])) < 0.005


def test_summaries(run, site):
    bins = pl.read_csv(site.correlation_summary_path("bins"))
    hist = pl.read_csv(site.correlation_summary_path("hist"))
    summary = pl.read_csv(site.correlation_summary_path("summary"))
    assert set(summary["channel"]) == {0, 1}
    assert bins["dt_s"].max() < site.correlation.max_dt_s
    # Every pair in the bins is counted once in the histogram.
    n_bins = bins.group_by("channel").agg(pl.col("n").sum()).sort("channel")
    n_hist = hist.group_by("channel").agg(pl.col("count").sum()).sort("channel")
    assert n_bins["n"].to_list() == n_hist["count"].to_list()
    # The pulses never decorrelate.
    for level in site.correlation.levels:
        assert summary[level_label(level)].is_infinite().all()
    assert (summary["rho_first"] > 0.9).all()


def test_decorrelation_time():
    dt = np.array([1.0, 2.0, 3.0, 4.0])
    rho = np.array([0.95, 0.9, 0.8, 0.7])
    assert decorrelation_time(dt, rho, 0.85) == pytest.approx(2.5)
    # Below the level already at the first bin: interpolated from rho = 1
    # at dt = 0.
    assert decorrelation_time(dt, rho, 0.975) == pytest.approx(0.5)
    assert decorrelation_time(dt, rho, 0.5) == np.inf


def test_binned_and_histogram_agree():
    config = CorrelationConfig(bin_s=2.0)
    rng = np.random.default_rng(0)
    dt = rng.uniform(0, 10, 1000)
    rho = rng.uniform(-0.1, 1.0, 1000)
    b = binned(rho, dt, config)
    h = histogram(rho, dt, config)
    assert b["dt_s"].to_list() == [1.0, 3.0, 5.0, 7.0, 9.0]
    assert b["n"].sum() == h["count"].sum() == 1000
    assert h["rho"].min() == 0.0


def test_stacked_image():
    t0 = np.datetime64("2023-07-06T15:00:00", "us")
    time = t0 + (np.arange(5) * 1.5e6).astype("timedelta64[us]")
    corr = np.full((5, 5), 0.5)
    np.fill_diagonal(corr, 1.0)
    corr[2] = corr[:, 2] = np.nan
    grid, t, image = stacked(corr, time, max_dt=3.0, n_grid=7)
    assert t.tolist() == [0.0, 1.5, 3.0, 4.5, 6.0]
    assert np.all(np.isnan(image[2]))
    np.testing.assert_allclose(image[0, [0, 3]], [1.0, 0.5])
    # The last strike has no strikes after it.
    assert np.all(np.isnan(image[4]))


def test_report(run, site):
    workflow.Workflow(site).report()
    figs = site.report_dir / "figures"
    assert (figs / "E1" / "correlation.png").exists()
    assert (figs / "E1" / "correlation_density.png").exists()
    assert not (figs / "correlation_density.png").exists()
    tex = (site.report_dir / "tables" / "E1_correlation.tex").read_text()
    assert "$T_{0.9}$ (s)" in tex and "> 60" in tex
    doc = docx.Document(str(site.report_dir / "report.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert "Correlation between pulses" not in texts
    assert any("Distribution of ρ" in t and "event E1" in t for t in texts)
    table = next(
        t for t in doc.tables if any("T0.9" in c.text for c in t.rows[0].cells)
    )
    assert [c.text for c in table.rows[1].cells][-2:] == ["> 60", "> 60"]


def test_strike_rows_leave_pauses_blank():
    from soundscape.report import ROW_HALF_S, strike_rows

    t = np.array([0.0, 1.5, 3.0, 300.0, 301.5])
    image = np.arange(10.0).reshape(5, 2)
    edges, rows = strike_rows(t, image)
    # One blank row for the pause between 3 s and 300 s.
    assert len(rows) == 6 and len(edges) == 7
    assert np.all(np.isnan(rows[3]))
    np.testing.assert_allclose(
        edges, [-0.75, 0.75, 2.25, 3.0 + ROW_HALF_S, 300 - ROW_HALF_S, 300.75, 302.25]
    )


def test_density_median_is_interpolated():
    from soundscape.report import correlation_density

    # 40% of pairs in the 0.50-0.52 bin and 60% in 0.52-0.54: the median
    # lies a sixth of the way into the second bin.
    h = pl.DataFrame({"dt_s": [1.0, 1.0], "rho": [0.50, 0.52], "count": [40, 60]})
    edges, percent, median = correlation_density(h, bin_s=2.0, max_dt=4.0)
    assert edges.tolist() == [0.0, 2.0, 4.0]
    assert median[0] == pytest.approx(0.52 + 0.02 / 6)
    assert np.isnan(median[1])
    assert np.nansum(percent[:, 0]) == pytest.approx(100.0)
    assert percent[25:27, 0].tolist() == pytest.approx([40.0, 60.0])


def test_compare_correlation_vs_range(run, site, tmp_path):
    from soundscape.compare import compare_sites

    out = tmp_path / "compare"
    compare_sites([site], out)
    assert (out / "figures" / "correlation_vs_range.png").exists()
    pts = pl.read_csv(out / "tables" / "correlation_vs_range.csv")
    # One row per event and sensor, on the detection channel.
    assert len(pts) == 1 and pts["channel"].to_list() == [0]
    assert {"rho_first", "rho_1min", "rho_5min", "range_km"} <= set(pts.columns)
    # max_dt_s is 60 s in this fixture, so the bin nearest 5 min is the last.
    assert pts["rho_5min"][0] > 0.9
    doc = docx.Document(str(out / "compare.docx"))
    assert any(
        p.text == "Correlation between pulses versus range" for p in doc.paragraphs
    )
